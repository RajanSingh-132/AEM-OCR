"""Word loading (PRD 5.4): .docx text and tables, plus OCR of pasted images."""

import io
import logging
import os
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document

from OCRAI.ocr import ocr_image
from OCRAI.pdf_loader import MIN_REAL_WORDS, real_words

log = logging.getLogger(__name__)

MEDIA_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
MIN_MEDIA_BYTES = 1_500   # smaller images are icons
MAX_OCR_IMAGES = 5
SOFFICE_TIMEOUT_S = 120
SOFFICE_NAMES = ("soffice", "soffice.exe", "libreoffice")
SOFFICE_PATHS = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
    "/usr/lib/libreoffice/program/soffice",
    "/opt/libreoffice/program/soffice",
    "/snap/bin/libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


def find_soffice() -> str | None:
    for name in SOFFICE_NAMES:
        if path := shutil.which(name):
            return path
    return next((path for path in SOFFICE_PATHS if os.path.isfile(path)), None)


def _convert_doc(data: bytes, filename: str) -> bytes:
    soffice = find_soffice()
    fallback = f"Could not convert '{filename}' (legacy .doc). Save it as .docx or .pdf and re-upload."
    if not soffice:
        raise ValueError(f"{fallback} (LibreOffice is not installed on the server.)")
    with tempfile.TemporaryDirectory() as tmp:
        source = Path(tmp) / "input.doc"
        source.write_bytes(data)
        try:
            subprocess.run(
                [soffice, "--headless", "--nologo", "--nolockcheck", "--nodefault",
                 "--nofirststartwizard", "--convert-to", "docx", "--outdir", tmp, str(source)],
                capture_output=True, timeout=SOFFICE_TIMEOUT_S, check=True,
            )
            return (Path(tmp) / "input.docx").read_bytes()
        except (subprocess.SubprocessError, OSError) as exc:
            log.error("  LibreOffice conversion failed: %s", exc)
            raise ValueError(fallback) from exc


def to_docx_bytes(data: bytes, filename: str) -> bytes:
    if not data:
        raise ValueError(f"Uploaded Word file '{filename}' is empty.")
    if Path(filename).suffix.lower() == ".docx" or data.startswith(b"PK"):
        return data
    log.info("  Legacy .doc: converting with LibreOffice")
    return _convert_doc(data, filename)


def docx_text(docx: bytes) -> str:
    document = Document(io.BytesIO(docx))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                text = " ".join(cell.text.split())
                # Merged cells repeat the same cell; keep it once.
                if text and (not cells or cells[-1] != text):
                    cells.append(text)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def docx_images(docx: bytes) -> list[tuple[str, bytes]]:
    """Embedded images (word/media/*), largest first, icons skipped."""
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        infos = [
            info for info in archive.infolist()
            if info.filename.startswith("word/media/")
            and Path(info.filename).suffix.lower() in MEDIA_EXTENSIONS
            and info.file_size >= MIN_MEDIA_BYTES
        ]
        infos.sort(key=lambda info: info.file_size, reverse=True)
        return [(Path(info.filename).name, archive.read(info)) for info in infos]


def load_word(data: bytes, filename: str) -> list[str]:
    docx = to_docx_bytes(data, filename)
    body = docx_text(docx).strip()
    images = docx_images(docx)
    word_count = len(real_words(body))
    log.info("  Word: body %d chars (%d words), %d embedded image(s)", len(body), word_count, len(images))

    ocr_text = ""
    if images and word_count < MIN_REAL_WORDS:
        texts = [ocr_image(image, name) for name, image in images[:MAX_OCR_IMAGES]]
        ocr_text = "\n".join(text for text in texts if text)
        log.info("  Word: OCR of %d image(s) -> %d chars", min(len(images), MAX_OCR_IMAGES), len(ocr_text))

    final = ocr_text if len(ocr_text) >= len(body) else body
    if not final:
        raise ValueError(
            f"No readable text found in '{filename}'. If the Word file has a pasted order image, "
            "ensure the image is embedded and vision OCR is configured; or export to PDF/image and re-upload."
        )
    return [final]
